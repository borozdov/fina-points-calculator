"""
Excel-версия калькулятора: excel/fina-points-calculator.xlsx.

Базовые времена берутся из js/data/constants.js, поэтому после смены редакции таблицы
достаточно перезапустить скрипт: python3 scripts/build-excel.py
Формулы Excel пересчитает сам при открытии файла.
"""
import re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'js/data/constants.js'
OUT = sys.argv[1] if len(sys.argv) > 1 else ROOT / 'excel/fina-points-calculator.xlsx'

text = open(SRC, encoding='utf8').read()
src_url = re.search(r"source:\s*'([^']+)'", text).group(1)
editions = {p: re.search(p + r":\s*\{\s*edition:\s*'([^']+)',\s*from:\s*'([\d-]+)',\s*to:\s*'([\d-]+)'", text).groups() for p in ('SCM', 'LCM')}
bt_block = text[text.index('export const BT'):text.index('export const RU')]

def parse_gender(pool, gender):
    pool_part = bt_block[bt_block.index(f'"{pool}"'):]
    seg = pool_part[pool_part.index(f'"{gender}"'):]
    seg = seg[seg.index('{') + 1:seg.index('}')]
    return dict((k, float(v)) for k, v in re.findall(r'"([^"]+)":\s*([\d.]+)', seg))

STYLE = {'Freestyle': 'вольный стиль', 'Backstroke': 'на спине', 'Breaststroke': 'брасс',
         'Butterfly': 'баттерфляй', 'Medley': 'комплекс'}
ORDER = ['50m Freestyle', '100m Freestyle', '200m Freestyle', '400m Freestyle', '800m Freestyle', '1500m Freestyle',
         '50m Backstroke', '100m Backstroke', '200m Backstroke', '50m Breaststroke', '100m Breaststroke', '200m Breaststroke',
         '50m Butterfly', '100m Butterfly', '200m Butterfly', '100m Medley', '200m Medley', '400m Medley']
def ru(ev):
    d, s = ev.split('m ')
    return f'{d} м {STYLE[s]}'

POOLS = [('SCM', '25 м'), ('LCM', '50 м')]
GENDERS = [('Men', 'Мужчины'), ('Women', 'Женщины')]

F = 'Arial'
font = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
title_font = Font(name=F, size=14, bold=True)
head_font = Font(name=F, size=10, bold=True, color='FFFFFF')
head_fill = PatternFill('solid', fgColor='000000')
input_fill = PatternFill('solid', fgColor='F2F2F2')
muted = Font(name=F, size=9, color='666666')
hair = Side(style='thin', color='BFBFBF')
box = Border(left=hair, right=hair, top=hair, bottom=hair)
center = Alignment(horizontal='center', vertical='center')

wb = Workbook()

# --- Базовые времена ---
base = wb.active
base.title = 'Базовые времена'
headers = ['Бассейн', 'Пол', 'Дистанция', 'Базовое время, с', 'Базовое время']
base.append(headers)
rows = 0
for p, pru in POOLS:
    for g, gru in GENDERS:
        data = parse_gender(p, g)
        for ev in ORDER:
            if ev in data:
                rows += 1
                r = rows + 1
                base.append([pru, gru, ru(ev), data[ev], None])
assert rows == 70, rows

def fmt_time(c):
    """Секунды → текст м:сс.сс без TEXT(), чтобы не зависеть от локали Excel."""
    cs = f'ROUND({c}*100,0)'
    m = f'INT({cs}/6000)'
    s = f'INT(MOD({cs},6000)/100)'
    h = f'MOD({cs},100)'
    return f'IF({m}>0,{m}&":"&RIGHT("0"&{s},2),{s})&"."&RIGHT("0"&{h},2)'

for r in range(2, rows + 2):
    base[f'E{r}'] = '=' + fmt_time(f'D{r}')
    base[f'D{r}'].number_format = '0.00'
    base[f'E{r}'].alignment = Alignment(horizontal='right')
note_r = rows + 3
base[f'A{note_r}'] = f'Источник: World Aquatics Points — Base Times SCM and LCM. Редакции: {editions["SCM"][0]} действует {editions["SCM"][1]} — {editions["SCM"][2]}; {editions["LCM"][0]} действует {editions["LCM"][1]} — {editions["LCM"][2]}.'
base[f'A{note_r + 1}'] = src_url
base[f'A{note_r + 2}'] = 'Когда редакция истечёт, замените числа в столбце D — остальные листы пересчитаются сами.'
for c in range(note_r, note_r + 3):
    base[f'A{c}'].font = muted
BASE_RANGE = f"'Базовые времена'!$D$2:$D${rows + 1}"
POOL_R = f"'Базовые времена'!$A$2:$A${rows + 1}"
G_R = f"'Базовые времена'!$B$2:$B${rows + 1}"
EV_R = f"'Базовые времена'!$C$2:$C${rows + 1}"

def base_lookup(pool, gender, ev):
    return f'SUMIFS({BASE_RANGE},{POOL_R},{pool},{G_R},{gender},{EV_R},{ev})'

def parse_time(c):
    """Время из ячейки: число секунд, время Excel (доля суток) или текст «1:02.35» / «28,45»."""
    s = f'SUBSTITUTE(TRIM({c}),",",".")'
    return (f'IF(ISNUMBER({c}),IF({c}<1,{c}*86400,{c}),'
            f'IFERROR(IF(ISNUMBER(FIND(":",{s})),'
            f'VALUE(LEFT({s},FIND(":",{s})-1))*60+_xlfn.NUMBERVALUE(MID({s},FIND(":",{s})+1,99),"."),'
            f'_xlfn.NUMBERVALUE({s},".")),0))')

# Списки для выпадающих меню
lists = wb.create_sheet('Списки')
lists['A1'] = 'Бассейн'; lists['B1'] = 'Пол'; lists['C1'] = 'Дистанция'
lists['D1'] = 'Минуты'; lists['E1'] = 'Секунды'; lists['F1'] = 'Сотые'; lists['G1'] = 'Очки'
for i, (_, pru) in enumerate(POOLS): lists.cell(i + 2, 1, pru)
for i, (_, gru) in enumerate(GENDERS): lists.cell(i + 2, 2, gru)
for i, ev in enumerate(ORDER): lists.cell(i + 2, 3, ru(ev))
MINUTES, SECONDS, HUNDREDTHS = range(0, 31), range(0, 60), range(0, 100)
POINTS = range(1000, 99, -10)
for col, values in ((4, MINUTES), (5, SECONDS), (6, HUNDREDTHS), (7, POINTS)):
    for i, v in enumerate(values): lists.cell(i + 2, col, v)
lists.sheet_state = 'hidden'
def dv_list(col, n):
    return lambda: DataValidation(type='list', formula1=f"='Списки'!${col}$2:${col}${n + 1}", allow_blank=True)
dv_pool, dv_g, dv_ev = dv_list('A', 2), dv_list('B', 2), dv_list('C', len(ORDER))
dv_min, dv_sec, dv_hun, dv_pts = dv_list('D', len(MINUTES)), dv_list('E', len(SECONDS)), dv_list('F', len(HUNDREDTHS)), dv_list('G', len(POINTS))

def slowest_time(b, pts):
    """Самое медленное время до сотых, которое ещё даёт столько очков. Эпсилон гасит хвосты плавающей точки."""
    return f'ROUNDDOWN({b}/({pts}/1000)^(1/3)+1E-9,2)'

def field(ws, r, label, value, dv, fmt=None):
    ws[f'A{r}'] = label; ws[f'A{r}'].font = bold
    ws[f'B{r}'] = value; ws[f'B{r}'].fill = input_fill; ws[f'B{r}'].border = box
    if fmt: ws[f'B{r}'].number_format = fmt
    v = dv(); v.add(f'B{r}'); ws.add_data_validation(v)

# --- Калькулятор: время из трёх списков ---
calc = wb.create_sheet('Калькулятор', 0)
calc['A1'] = 'Калькулятор очков FINA (World Aquatics Points)'
calc['A1'].font = title_font
calc['A2'] = 'Всё выбирается из списков в серых ячейках. Можно и напечатать число в ячейку.'
calc['A2'].font = muted
field(calc, 4, 'Бассейн', '25 м', dv_pool)
field(calc, 5, 'Пол', 'Мужчины', dv_g)
field(calc, 6, 'Дистанция', '100 м вольный стиль', dv_ev)
field(calc, 7, 'Минуты', 0, dv_min)
field(calc, 8, 'Секунды', 50, dv_sec, '00')
field(calc, 9, 'Сотые', 12, dv_hun, '00')

calc['A11'] = 'Время'; calc['B11'] = '=IF(B13="","",' + fmt_time('B13') + ')'
calc['A12'] = 'Базовое время, с'; calc['B12'] = '=' + base_lookup('$B$4', '$B$5', '$B$6')
calc['A13'] = 'Время, с'; calc['B13'] = '=IF(N(B7)*60+N(B8)+N(B9)/100<=0,"",N(B7)*60+N(B8)+N(B9)/100)'
calc['A14'] = 'Очки FINA'; calc['A14'].font = bold
calc['B14'] = '=IF(OR(B12=0,B13=""),"",ROUNDDOWN(1000*(B12/B13)^3,0))'
calc['B14'].font = Font(name=F, size=16, bold=True)
calc['C12'] = '=IF(B12=0,"Такой дистанции нет в этом бассейне","")'
calc['C12'].font = muted
calc['B12'].number_format = '0.00'; calc['B13'].number_format = '0.00'
for r in (11, 12, 13): calc[f'B{r}'].alignment = Alignment(horizontal='right')

calc['A16'] = 'Обратный расчёт'; calc['A16'].font = bold
field(calc, 17, 'Нужные очки', 700, dv_pts)
calc['A18'] = 'Нужное время'; calc['A18'].font = bold
calc['C18'] = '=IF(OR(B12=0,N(B17)<=0),"",' + slowest_time('B12', 'B17') + ')'
calc['B18'] = '=IF(C18="","",' + fmt_time('C18') + ')'
calc['B18'].font = Font(name=F, size=16, bold=True); calc['B18'].alignment = Alignment(horizontal='right')
calc['C18'].font = muted; calc['C18'].number_format = '0.00'

calc['A20'] = 'Формула: очки = 1000 × (базовое время / время)³, дробная часть отбрасывается.'
calc['A21'] = 'Сайт: https://fina.borozdov.ru — тот же расчёт.'
for c in ('A20', 'A21'): calc[c].font = muted

# --- Протокол ---
prot = wb.create_sheet('Протокол', 1)
ph = ['Спортсмен', 'Бассейн', 'Пол', 'Дистанция', 'Время', 'Очки FINA', 'Базовое время, с', 'Время, с']
prot.append(ph)
example = [('Иванов Иван', '25 м', 'Мужчины', '100 м вольный стиль', '50.12'),
           ('Петрова Анна', '50 м', 'Женщины', '200 м комплекс', '2:21.40')]
N = 500
for r in range(2, N + 2):
    if r - 2 < len(example):
        for j, v in enumerate(example[r - 2]):
            prot.cell(r, j + 1, v)
    for col in 'ABCDE':
        prot[f'{col}{r}'].fill = input_fill
        prot[f'{col}{r}'].border = box
    prot[f'E{r}'].number_format = '@'
    prot[f'G{r}'] = f'=IF(D{r}="","",' + base_lookup(f'B{r}', f'C{r}', f'D{r}') + ')'
    prot[f'H{r}'] = f'=IF(E{r}="","",' + parse_time(f'E{r}') + ')'
    prot[f'F{r}'] = f'=IF(OR(G{r}="",H{r}="",N(G{r})=0,N(H{r})<=0),"",ROUNDDOWN(1000*(G{r}/H{r})^3,0))'
    prot[f'F{r}'].font = bold
    prot[f'G{r}'].number_format = '0.00'; prot[f'H{r}'].number_format = '0.00'
    for col in 'GH':
        prot[f'{col}{r}'].font = muted
for dv, col in ((dv_pool(), 'B'), (dv_g(), 'C'), (dv_ev(), 'D')):
    dv.add(f'{col}2:{col}{N + 1}'); prot.add_data_validation(dv)
prot['F1'].comment = Comment('Пусто, если не заполнена строка или дистанции нет в этом бассейне (например, 100 м комплекс в 50 м).', 'fina.borozdov.ru')
prot.freeze_panes = 'A2'
prot.auto_filter.ref = f'A1:H{N + 1}'

# Шапки и шрифты
for ws in (base, prot):
    for c in ws[1]:
        c.font = head_font; c.fill = head_fill; c.alignment = center
for ws in (calc, base, prot):
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font():  # не тронутые стилем
                c.font = font
            elif c.font.name != F:
                c.font = Font(name=F, size=c.font.size, bold=c.font.bold, color=c.font.color)
for ws, widths in ((calc, [20, 22, 40]), (base, [10, 12, 24, 18, 16]), (prot, [28, 10, 12, 24, 12, 12, 18, 12])):
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w

# По частоте использования: калькулятор, протокол, справочник
wb._sheets = [calc, prot, base, lists]
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT)
